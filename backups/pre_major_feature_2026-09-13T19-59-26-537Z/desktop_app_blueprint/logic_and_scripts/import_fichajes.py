#!/usr/bin/env python3
"""
TPLH Forecast — Skrypt Importu Logowań MAPAL (Fichajes)
Obsługuje automatyczne czyszczenie (pominięcie wiersza 7),
filtrację dla kawiarni 108120 SBX Warszawa Janki,
kolumnę Computable Time oraz aktualizację bazy 'Labor Actuals DB'.

Użycie:
    python import_fichajes.py [sciezka_do_raportu.xls]
"""

import sys
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import xlrd
from datetime import datetime

DEFAULT_REPORT = "amrestpl_8cc166cf-8958-4229-9ca5-eb17796d517c_11920261852501735_Fichajes.xls"
PRODUCTION_FILE = "TPLH_Forecast_v1.0_CleanProduction.xlsx"

def import_report(report_path=None):
    if not report_path:
        report_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_REPORT
        
    if not os.path.exists(report_path):
        print(f"❌ Błąd: Nie znaleziono pliku raportu: {report_path}")
        return False

    print(f"📂 Otwieranie raportu: {report_path}")
    wb_xls = xlrd.open_workbook(report_path)
    sheet_xls = wb_xls.sheet_by_index(0)

    if sheet_xls.nrows < 8:
        print("❌ Błąd: Plik nie zawiera danych logowań (mniej niż 8 wierszy).")
        return False

    print(f"📖 Otwieranie pliku produkcyjnego: {PRODUCTION_FILE}")
    wb_main = openpyxl.load_workbook(PRODUCTION_FILE)
    
    if "Calendar Engine" not in wb_main.sheetnames:
        print("❌ Błąd: Brak arkusza 'Calendar Engine' w pliku produkcyjnym.")
        return False
        
    cal = wb_main["Calendar Engine"]
    cal_map = {}
    for r in range(2, cal.max_row + 1):
        d_val = cal.cell(r, 1).value
        if isinstance(d_val, datetime):
            d_str = d_val.strftime("%Y-%m-%d")
            cal_map[d_str] = {
                "rok": cal.cell(r, 5).value,
                "msc": cal.cell(r, 6).value,
                "dzien": cal.cell(r, 7).value,
                "tydzien": cal.cell(r, 9).value,
                "klucz": cal.cell(r, 10).value
            }

    if "Labor Actuals DB" not in wb_main.sheetnames:
        ws_db = wb_main.create_sheet(title="Labor Actuals DB")
        headers = [
            "ID", "Data", "Rok", "Miesiąc", "Tydzień", "Klucz_Tygodnia", 
            "Dzień", "Pracownik", "Stanowisko", "Wymiar_Etatu", 
            "Godziny", "Kawiarnia_Kod", "Kawiarnia_Nazwa", "Data_Importu"
        ]
        ws_db.append(headers)
        header_fill = PatternFill(start_color="1E3932", end_color="1E3932", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for c in range(1, len(headers) + 1):
            cell = ws_db.cell(1, c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
    else:
        ws_db = wb_main["Labor Actuals DB"]

    # Zbieranie nowych wpisów dla Janki
    import_time = datetime.now().strftime("%Y-%m-%d %H:%M")
    new_records = []
    dates_in_import = set()

    # Omijamy wiersz 7 (index 6), czytamy od wiersza 8 (index 7)
    for r in range(7, sheet_xls.nrows):
        row = sheet_xls.row_values(r)
        unit = str(row[15])
        if "108120" in unit or "Janki" in unit:
            b_day = row[6]
            d = xlrd.xldate_as_datetime(b_day, wb_xls.datemode)
            d_str = d.strftime("%Y-%m-%d")
            dates_in_import.add(d_str)
            
            cal_info = cal_map.get(d_str, {})
            emp = str(row[3]).strip()
            cat = str(row[4]).strip()
            contract = str(row[5]).strip()
            c_time = round(float(row[9]), 2)
            unit_code = str(row[14]).strip()

            new_records.append([
                d_str,
                cal_info.get("rok", d.year),
                cal_info.get("msc", ""),
                cal_info.get("tydzien", ""),
                cal_info.get("klucz", ""),
                cal_info.get("dzien", ""),
                emp,
                cat,
                contract,
                c_time,
                unit_code,
                unit,
                import_time
            ])

    if not new_records:
        print("⚠️ Brak rekordów dla kawiarni 108120 SBX Janki w pliku.")
        return False

    min_date = min(dates_in_import)
    max_date = max(dates_in_import)
    print(f"📊 Znaleziono {len(new_records)} wpisów Janki w zakresie: {min_date} do {max_date}")

    # Idempotencja: Usuwamy istniejące rekordy w DB z tego samego zakresu dat
    existing_rows = list(ws_db.iter_rows(min_row=2, values_only=True))
    kept_rows = []
    for r in existing_rows:
        row_date = str(r[1]) if len(r) > 1 else ""
        if row_date < min_date or row_date > max_date:
            kept_rows.append(r)

    # Przebudowa zawartości DB
    ws_db.delete_rows(2, ws_db.max_row)
    for idx, r in enumerate(kept_rows, start=1):
        ws_db.append([idx] + list(r[1:]))

    current_id = len(kept_rows)
    for r in new_records:
        current_id += 1
        ws_db.append([current_id] + r)

    # Formatowanie kolumn
    for col in ws_db.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_db.column_dimensions[col_letter].width = max(max_len + 3, 11)

    ws_db.auto_filter.ref = f"A1:N{ws_db.max_row}"

    # Połączenie z Dashboardem (J12:J17)
    dash = wb_main["Executive Dashboard"]
    for idx, row_idx in enumerate(range(12, 18), start=1):
        w_key = f"_W{idx}"
        dash.cell(row_idx, 10).value = (
            f"=IF(SUMIFS('Labor Actuals DB'!$K:$K, 'Labor Actuals DB'!$F:$F, "
            f"$C$5 & \"_\" & $C$6 & \"{w_key}\")>0, "
            f"ROUND(SUMIFS('Labor Actuals DB'!$K:$K, 'Labor Actuals DB'!$F:$F, "
            f"$C$5 & \"_\" & $C$6 & \"{w_key}\"), 1), \"\")"
        )

    # Utrzymanie trybu jednoarkuszowego (tylko Executive Dashboard widoczny)
    for name in wb_main.sheetnames:
        if name != "Executive Dashboard":
            wb_main[name].sheet_state = "hidden"
        else:
            wb_main[name].sheet_state = "visible"

    wb_main.save(PRODUCTION_FILE)
    print(f"✅ Sukces! Zaktualizowano bazę danych ({ws_db.max_row - 1} rekordów) i plik produkcyjny {PRODUCTION_FILE}.")
    return True

if __name__ == "__main__":
    import_report()
